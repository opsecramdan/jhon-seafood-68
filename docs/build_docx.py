import re, json
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.enum.text import WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

REPORT_MD = "/Users/macbook/ramdan/kampus/SEMESTER_5/CLOUD COMPUTING/jhon-seafood-68/docs/laporan-draft.md"
OUT_DOCX = "/Users/macbook/ramdan/kampus/SEMESTER_5/CLOUD COMPUTING/jhon-seafood-68/docs/CC-Kel01-05SIFE001.docx"

with open(REPORT_MD, encoding="utf-8") as f:
    lines = f.read().split("\n")

def clean_inline(text):
    text = re.sub(r'\*\*(.*?)\*\*', r'\1', text)
    text = re.sub(r'`([^`]*)`', r'\1', text)
    text = text.replace('&mdash;', '-').replace('&middot;', '.')
    return text.strip()

blocks = []
i = 0
n = len(lines)
while i < n:
    line = lines[i]
    stripped = line.strip()
    if stripped == '' or stripped == '---' or stripped.startswith('>'):
        i += 1
        continue
    m = re.match(r'^(#{1,6})\s+(.*)', stripped)
    if m:
        level = len(m.group(1))
        text = clean_inline(m.group(2))
        blocks.append({"type": "heading", "level": level, "text": text})
        i += 1
        continue
    if stripped.startswith('```'):
        i += 1
        code_lines = []
        while i < n and not lines[i].strip().startswith('```'):
            code_lines.append(lines[i])
            i += 1
        i += 1
        for cl in code_lines:
            blocks.append({"type": "code", "text": cl if cl.strip() else " "})
        continue
    if stripped.startswith('|'):
        table_lines = []
        while i < n and lines[i].strip().startswith('|'):
            table_lines.append(lines[i].strip())
            i += 1
        rows = []
        for tl in table_lines:
            cells = [c.strip() for c in tl.strip('|').split('|')]
            if all(re.match(r'^:?-+:?$', c) for c in cells):
                continue
            rows.append([clean_inline(c) for c in cells])
        if rows:
            blocks.append({"type": "table", "header": rows[0], "rows": rows[1:]})
        continue
    m = re.match(r'^\d+\.\s+(.*)', stripped)
    if m:
        items = []
        while i < n:
            s = lines[i].strip()
            mm = re.match(r'^\d+\.\s+(.*)', s)
            if mm:
                items.append(clean_inline(mm.group(1)))
                i += 1
            elif lines[i].startswith('   ') and s != '':
                if items:
                    items[-1] += ' ' + clean_inline(s)
                i += 1
            else:
                break
        blocks.append({"type": "numbered_list", "items": items})
        continue
    m = re.match(r'^[-*]\s+(.*)', stripped)
    if m:
        items = []
        while i < n:
            s = lines[i].strip()
            mm = re.match(r'^[-*]\s+(.*)', s)
            if mm:
                items.append(clean_inline(mm.group(1)))
                i += 1
            elif lines[i].startswith('  ') and s != '' and not re.match(r'^[-*]\s+', s) and not re.match(r'^#{1,6}\s', s):
                if items:
                    items[-1] += ' ' + clean_inline(s)
                i += 1
            else:
                break
        blocks.append({"type": "bullet_list", "items": items})
        continue
    para_lines = [stripped]
    i += 1
    while i < n and lines[i].strip() != '' and not re.match(r'^(#{1,6}\s|\d+\.\s|[-*]\s|\||```|>)', lines[i].strip()):
        para_lines.append(lines[i].strip())
        i += 1
    text = clean_inline(' '.join(para_lines))
    blocks.append({"type": "paragraph", "text": text})

# Drop first two generic headings (draft title placeholders)
body_blocks = blocks[2:]

# ---------------- Build DOCX ----------------
doc = Document()
section = doc.sections[0]
section.page_width = Mm(210)
section.page_height = Mm(297)
section.top_margin = Mm(25)
section.bottom_margin = Mm(25)
section.left_margin = Mm(30)
section.right_margin = Mm(25)

normal = doc.styles['Normal']
normal.font.name = 'Times New Roman'
normal.font.size = Pt(12)
pf = normal.paragraph_format
pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
pf.space_after = Pt(6)
rpr = normal.element.get_or_add_rPr()
rFonts = rpr.find(qn('w:rFonts'))
if rFonts is None:
    rFonts = OxmlElement('w:rFonts')
    rpr.append(rFonts)
rFonts.set(qn('w:eastAsia'), 'Times New Roman')

heading_sizes = {1: 16, 2: 14, 3: 13}
for lvl, sz in heading_sizes.items():
    st = doc.styles[f'Heading {lvl}']
    st.font.name = 'Times New Roman'
    st.font.size = Pt(sz)
    st.font.bold = True
    st.font.color.rgb = RGBColor(0, 0, 0)
    st.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    st.paragraph_format.space_before = Pt(12)
    st.paragraph_format.space_after = Pt(6)

from docx.enum.style import WD_STYLE_TYPE
code_style = doc.styles.add_style('CodeBlock', WD_STYLE_TYPE.PARAGRAPH)
code_style.base_style = doc.styles['Normal']
code_style.font.name = 'Consolas'
code_style.font.size = Pt(9)
code_style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
code_style.paragraph_format.space_after = Pt(0)

def add_table(doc, header, rows):
    ncols = len(header) if header else (len(rows[0]) if rows else 1)
    table = doc.add_table(rows=0, cols=ncols)
    table.style = 'Table Grid'
    if header:
        cells = table.add_row().cells
        for idx, text in enumerate(header):
            cells[idx].text = str(text)
            for para in cells[idx].paragraphs:
                para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
                for run in para.runs:
                    run.bold = True
                    run.font.size = Pt(10)
    for row in rows:
        cells = table.add_row().cells
        for idx, text in enumerate(row):
            if idx < len(cells):
                cells[idx].text = str(text)
                for para in cells[idx].paragraphs:
                    para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
                    for run in para.runs:
                        run.font.size = Pt(10)

# --- Title page ---
from docx.enum.text import WD_ALIGN_PARAGRAPH

def centered(doc, text, size=12, bold=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(size)
    return p

for _ in range(3):
    doc.add_paragraph()

centered(doc, "LAPORAN TUGAS PRAKTIKUM CLOUD COMPUTING", size=16, bold=True)
centered(doc, "Migrasi Sistem UMKM ke Resources Cloud", size=14, bold=True)
centered(doc, "Studi Kasus: Jhon Seafood 68", size=14, bold=True)
doc.add_paragraph()
doc.add_paragraph()
centered(doc, "Kelompok: K-1", size=12, bold=True)
doc.add_paragraph()

add_table(doc, ["NIM", "Nama", "Peran"], [
    ["241011700786", "Muhammad Ramadan", "Cloud Architect (Ketua)"],
    ["241011700832", "Deka Ramdoni", "Infrastructure Engineer"],
    ["241011701033", "Gopriadi", "Security & Governance Officer"],
    ["241011701152", "Muhammad Daffa Cahyo Santoso", "Cost & Operations Analyst"],
])

for _ in range(3):
    doc.add_paragraph()

centered(doc, "Mata Kuliah: Cloud Computing", size=12)
centered(doc, "Dosen Pengampu: Dendi Sunardi, M.Kom.", size=12)
centered(doc, "Program Studi Sistem Informasi", size=12)
centered(doc, "Fakultas Ilmu Komputer", size=12)
centered(doc, "Universitas Pamulang", size=12)

doc.add_page_break()

# --- Body ---
heading_level_map = {2: 1, 3: 2, 4: 3}
for b in body_blocks:
    t = b["type"]
    if t == "heading":
        lvl = heading_level_map.get(b["level"], 3)
        doc.add_heading(b["text"], level=lvl)
    elif t == "paragraph":
        doc.add_paragraph(b["text"])
    elif t == "code":
        doc.add_paragraph(b["text"], style="CodeBlock")
    elif t == "bullet_list":
        for item in b["items"]:
            doc.add_paragraph(item, style="List Bullet")
    elif t == "numbered_list":
        for item in b["items"]:
            doc.add_paragraph(item, style="List Number")
    elif t == "table":
        add_table(doc, b["header"], b["rows"])
        doc.add_paragraph()

# --- Footer: page numbers ---
from docx.oxml.ns import qn as _qn

def add_field(paragraph, instr_text, result_text="1"):
    run = paragraph.add_run()
    fld_begin = OxmlElement('w:fldChar')
    fld_begin.set(_qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText')
    instr.set(_qn('xml:space'), 'preserve')
    instr.text = instr_text
    fld_sep = OxmlElement('w:fldChar')
    fld_sep.set(_qn('w:fldCharType'), 'separate')
    fld_text = OxmlElement('w:t')
    fld_text.text = result_text
    fld_end = OxmlElement('w:fldChar')
    fld_end.set(_qn('w:fldCharType'), 'end')
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_sep)
    run._r.append(fld_text)
    run._r.append(fld_end)

footer_para = doc.sections[0].footer.paragraphs[0]
footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer_para.add_run("Halaman ")
add_field(footer_para, " PAGE ", "1")
footer_para.add_run(" dari ")
add_field(footer_para, " NUMPAGES ", "1")

doc.save(OUT_DOCX)
print(json.dumps({"ok": True, "output": OUT_DOCX, "blocks": len(body_blocks)}))
